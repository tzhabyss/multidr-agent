"""Sandboxed local document reader for worker-scoped file Skills."""

from __future__ import annotations

import csv
import io
import json
import mimetypes
import os
from pathlib import Path

from .schemas import FetchedDocument, utc_now


DEFAULT_MAX_FILE_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_CHARS = 50_000
TEXT_EXTENSIONS = {
    ".cfg",
    ".csv",
    ".ini",
    ".log",
    ".markdown",
    ".md",
    ".py",
    ".rst",
    ".text",
    ".toml",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}


class LocalFileSecurityError(ValueError):
    """Raised when a local file is outside the configured read roots."""


class LocalFileReader:
    """Read local documents under explicit, process-local directory roots.

    Roots are configured with ``DEEPRESEARCH_FILE_ROOTS`` as an OS-path-list.
    When unset, only the current working directory is readable. The handler
    returns a FetchedDocument-shaped record so the existing extraction and
    citation pipeline can process local content without treating it as a web
    page.
    """

    def __init__(
        self,
        roots: list[str] | None = None,
        max_file_bytes: int | None = None,
    ) -> None:
        configured = roots
        if configured is None:
            raw = os.getenv("DEEPRESEARCH_FILE_ROOTS", "")
            configured = [item for item in raw.split(os.pathsep) if item.strip()]
        if not configured:
            configured = [os.getcwd()]
        self.roots = tuple(self._resolve_root(item) for item in configured)
        self.max_file_bytes = max_file_bytes or int(
            os.getenv("DEEPRESEARCH_FILE_MAX_BYTES", str(DEFAULT_MAX_FILE_BYTES))
        )

    @staticmethod
    def _resolve_root(value: str) -> Path:
        root = Path(value).expanduser().resolve()
        if not root.exists() or not root.is_dir():
            raise LocalFileSecurityError(f"file root is not a directory: {root}")
        return root

    def _resolve_path(self, value: str) -> Path:
        if not isinstance(value, str) or not value.strip():
            raise LocalFileSecurityError("path must be a non-empty string")
        requested = Path(value).expanduser()
        lexical = Path(os.path.abspath(requested))
        current = lexical
        while True:
            if current.is_symlink():
                raise LocalFileSecurityError("symlink paths are not allowed")
            if current.parent == current:
                break
            current = current.parent
        path = requested.resolve()
        if not any(path == root or root in path.parents for root in self.roots):
            allowed = ", ".join(str(root) for root in self.roots)
            raise LocalFileSecurityError(f"path is outside configured file roots: {allowed}")
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(str(path))
        if path.stat().st_size > self.max_file_bytes:
            raise LocalFileSecurityError(
                f"file exceeds the {self.max_file_bytes}-byte limit: {path}"
            )
        return path

    @staticmethod
    def _read_text(path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="replace")

    @staticmethod
    def _read_csv(path: Path) -> str:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        rows = list(csv.reader(io.StringIO(LocalFileReader._read_text(path)), delimiter=delimiter))
        return "\n".join("\t".join(row) for row in rows)

    @staticmethod
    def _read_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError(
                "PDF reading requires the optional 'documents' dependencies: "
                "pip install -e .[documents]"
            ) from exc
        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    @staticmethod
    def _read_docx(path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "DOCX reading requires the optional 'documents' dependencies: "
                "pip install -e .[documents]"
            ) from exc
        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def read(self, path: str, max_chars: int = DEFAULT_MAX_CHARS) -> dict:
        """Return a bounded local document record for evidence extraction."""

        if max_chars < 100 or max_chars > DEFAULT_MAX_CHARS:
            raise ValueError(f"max_chars must be between 100 and {DEFAULT_MAX_CHARS}")
        resolved = self._resolve_path(path)
        suffix = resolved.suffix.lower()
        if suffix in TEXT_EXTENSIONS:
            content = self._read_csv(resolved) if suffix in {".csv", ".tsv"} else self._read_text(resolved)
        elif suffix == ".json":
            content = json.dumps(json.loads(self._read_text(resolved)), ensure_ascii=False, indent=2)
        elif suffix == ".pdf":
            content = self._read_pdf(resolved)
        elif suffix == ".docx":
            content = self._read_docx(resolved)
        else:
            raise ValueError(
                "unsupported file type; supported types include text, JSON, CSV, PDF, and DOCX"
            )
        content = content[:max_chars]
        return FetchedDocument(
            requested_url=path,
            final_url=resolved.as_uri(),
            title=resolved.name,
            content=content,
            content_type=mimetypes.guess_type(resolved.name)[0] or "text/plain",
            http_status=200,
            fetched_at=utc_now(),
            redirect_chain=[],
        ).model_dump(mode="json")
