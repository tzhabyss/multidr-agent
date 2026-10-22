"""Opt-in REAL MongoDB + Milvus tests. No mongomock or in-memory vector substitute.

Storage tests deliberately use deterministic hash vectors (allow_degraded=True),
not semantic-quality metrics. The separately named E5 test forbids that fallback.
Each test owns a fresh disposable Mongo database and removes only that database.
"""
import os
import json
import socket
import subprocess
import sys
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from pathlib import Path
from threading import Event

import pytest

pytestmark = pytest.mark.skipif(os.environ.get('DEEPRESEARCH_RAG_INTEGRATION') != '1',
                                reason='set DEEPRESEARCH_RAG_INTEGRATION=1 for real dual-database tests')


@pytest.fixture(scope='module')
def vector_uri(tmp_path_factory):
    pytest.importorskip('pymilvus')
    configured = os.environ.get('RAG_TEST_MILVUS_URI')
    if configured:
        yield configured
        return
    from milvus_lite.server import Server
    with socket.socket() as sock:
        sock.bind(('127.0.0.1',0)); port=sock.getsockname()[1]
    server=Server(str(tmp_path_factory.mktemp('milvus')/'live.db'), address=f'127.0.0.1:{port}')
    assert server.init() and server.start()
    try:
        yield f'http://127.0.0.1:{port}'
    finally:
        server.stop()


@pytest.fixture
def store(vector_uri):
    from deepresearch.knowledge.dual_store import MongoMilvusStore
    from deepresearch.memory_system.embeddings import HashEmbeddingProvider
    class Counter(HashEmbeddingProvider):
        def __init__(self): super().__init__(32); self.calls=0
        def encode_documents(self, values):
            self.calls+=1
            return super().encode_documents(values)
    db=MongoMilvusStore(os.environ.get('RAG_TEST_MONGO_URI','mongodb://127.0.0.1:27028'),
        database='rag_test_'+uuid.uuid4().hex, milvus_uri=vector_uri,
        provider=Counter(), allow_degraded=True)
    try:
        yield db
    finally:
        # Test namespaces only; never touch configured user databases.
        names={v.get('vector_collection') for v in db.db.versions.find({})}
        names|={j.get('vector_collection') for j in db.db.jobs.find({})}
        for name in names-{None}:
            if not db.vectors.has_collection(collection_name=name):
                continue
            # Vector collections are shared by model; delete only this test's IDs.
            ids=[cid for j in db.db.jobs.find({}) if j.get('vector_collection')==name for cid in j.get('milvus_ids',[])]
            for i in range(0,len(ids),128):
                db.vectors.delete(collection_name=name,ids=ids[i:i+128])
        db.mongo.drop_database(db.db.name)
        db.close()


def seed(db):
    return db.ingest('旧版费用为100元。',uri='file:///old.md',title='制度',document_id='policy')


def test_upload_sha_precedes_parse_and_embedding(store,tmp_path,monkeypatch):
    path=tmp_path/'report.md';path.write_text('费用为128元。')
    result=store.ingest_file(path); calls=store.provider.calls
    monkeypatch.setattr(store,'_parse',lambda *args: (_ for _ in ()).throw(AssertionError('must skip parsing')))
    duplicate=store.ingest_file(path)
    assert duplicate['duplicate'] and duplicate['snapshot_id']==result['snapshot_id']
    assert store.provider.calls==calls and store.db.versions.count_documents({})==1


def test_identical_upload_alias_avoids_reencoding(store,tmp_path):
    a,b=tmp_path/'a.md',tmp_path/'b.md'
    a.write_text('同一份原始报告。');b.write_bytes(a.read_bytes())
    first=store.ingest_file(a); calls=store.provider.calls
    second=store.ingest_file(b)
    assert second['duplicate'] and second['snapshot_id']==first['snapshot_id']
    assert store.provider.calls==calls and len(store.db.uploads.find_one({})['aliases'])==2


def test_renamed_versions_use_explicit_logical_id(store,tmp_path):
    a,b=tmp_path/'2024.md',tmp_path/'2025.md';a.write_text('费用100元。');b.write_text('费用200元。')
    old=store.ingest_file(a,document_id='财务制度-1'); new=store.ingest_file(b,document_id='财务制度-1')
    assert store.snapshot(old['snapshot_id'],collection='library')['status']=='archived'
    assert store.snapshot(new['snapshot_id'],collection='library')['supersedes']==old['snapshot_id']
    assert store.db.heads.count_documents({'active':{'$ne':None}})==1


@pytest.mark.parametrize('stage',['snapshot_written','mongo_chunks_written','vector_batch_written','ready'])
def test_failure_before_publication_compensates_both_planes(store,stage):
    old=seed(store)
    def fail(current,job):
        if current==stage: raise OSError('injected write interruption')
    store._stage=fail
    with pytest.raises(OSError):
        store.ingest('新版费用200元。\n\n'*150,uri='file:///new.md',title='制度',document_id='policy')
    job=store.db.jobs.find_one({'status':'failed'})
    assert job and job['phase']=='compensated' and job['error']
    assert store.db.chunks.count_documents({'snapshot':job['snapshot_id']})==0
    assert store.snapshot(old['snapshot_id'],collection='library')['status']=='active'
    assert store.snapshot(job['snapshot_id'],collection='library')['status']=='pending'
    if job.get('vector_collection'):
        assert not store.vectors.get(collection_name=job['vector_collection'],ids=job['milvus_ids'][:128])


@pytest.mark.parametrize('bad',['nan','inf','zero','dim','string','bool','empty','short'])
def test_invalid_embedding_never_archives_old(store,bad):
    old=seed(store)
    from deepresearch.memory_system.embeddings import HashEmbeddingProvider
    class Invalid(HashEmbeddingProvider):
        @property
        def degraded(self):return False
        def encode_documents(self, values):
            vectors=super().encode_documents(values)
            if bad=='short':return vectors[:-1]
            if bad=='dim': vectors[0]=[1.,2.,3.]
            elif bad=='empty':vectors[0]=[]
            elif bad=='zero':vectors[0]=[0.]*self.dimensions
            else:vectors[0][0]={'nan':float('nan'),'inf':float('inf'),'string':'x','bool':True}[bad]
            return vectors
    store.provider=Invalid(32)
    with pytest.raises(ValueError):
        store.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
    assert store.snapshot(old['snapshot_id'],collection='library')['status']=='active'
    assert store.db.jobs.count_documents({'status':'failed'})==1


def test_published_head_survives_ambiguous_ack_and_reconcile(store):
    old=seed(store)
    def fail(stage,job):
        if stage=='head_published': raise TimeoutError('lost publish acknowledgement')
    store._stage=fail
    with pytest.raises(TimeoutError):store.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
    job=store.db.jobs.find_one({'status':'pending'})
    assert store.snapshot(job['snapshot_id'],collection='library')['status']=='active'
    store._stage=lambda *args:None
    recovered=store.reconcile()
    assert not recovered['errors'] and recovered['published']==1
    assert store.db.jobs.find_one({'_id':job['_id']})['status']=='completed'
    assert store.snapshot(old['snapshot_id'],collection='library')['status']=='archived'
    assert '200' in store.search('费用',mode='bm25')[0]['content']


def test_cleanup_failure_is_durable_and_reconciled(store):
    old=seed(store)
    def fail(stage,job):
        if stage=='vector_batch_written':raise OSError('write failure')
        if stage=='vectors_compensated':raise OSError('cleanup failure')
    store._stage=fail
    with pytest.raises(OSError):store.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
    job=store.db.jobs.find_one({'status':'cleanup_failed'})
    assert job and job['cleanup_error']
    store._stage=lambda *args:None
    assert store.reconcile()['compensated']==1
    assert store.db.jobs.find_one({'_id':job['_id']})['status']=='failed'
    assert store.snapshot(old['snapshot_id'],collection='library')['status']=='active'


def test_failed_upload_retry_preserves_byte_identity(store,tmp_path):
    path=tmp_path/'retry.md';path.write_text('费用100元。')
    store._stage=lambda stage,job: (_ for _ in ()).throw(OSError('interrupt')) if stage=='vector_batch_written' else None
    with pytest.raises(OSError):store.ingest_file(path,document_id='policy')
    failed=store.db.jobs.find_one({'status':'failed'})
    store._stage=lambda *args:None
    result=store.retry(failed['_id']); calls=store.provider.calls
    again=store.ingest_file(path,document_id='policy')
    assert again['duplicate'] and result['snapshot_id']==again['snapshot_id'] and store.provider.calls==calls
    assert store.db.jobs.find_one({'_id':result['job_id']})['retry_of']==failed['_id']


def test_live_reader_and_second_writer_are_fenced(store,vector_uri):
    from deepresearch.knowledge.dual_store import MongoMilvusStore,IngestionBusy
    old=seed(store); reached,release=Event(),Event()
    def pause(stage,job):
        if stage=='ready':reached.set();assert release.wait(10)
    store._stage=pause
    other=MongoMilvusStore(mongo_client=store.mongo,database=store.db.name,vector_client=store.vectors,
                           provider=store.provider,allow_degraded=True)
    with ThreadPoolExecutor(max_workers=1) as pool:
        task=pool.submit(store.ingest,'新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
        try:
            assert reached.wait(10)
            assert all('100' in other.search('费用',mode='bm25')[0]['content'] for _ in range(10))
            with pytest.raises(IngestionBusy):other.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
        finally:release.set()
        result=task.result()
    calls=store.provider.calls
    assert other.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')['duplicate']
    assert store.provider.calls==calls
    assert other.snapshot(old['snapshot_id'],collection='library')['status']=='archived'
    assert other.snapshot(result['snapshot_id'],collection='library')['status']=='active'


@pytest.mark.parametrize('stage',['vector_batch_written','head_published'])
def test_real_process_exit_and_durable_recovery(store,vector_uri,stage):
    from deepresearch.knowledge.dual_store import utcnow
    old=seed(store)
    script='''
import os,sys
from deepresearch.knowledge.dual_store import MongoMilvusStore
from deepresearch.memory_system.embeddings import HashEmbeddingProvider
s=MongoMilvusStore(sys.argv[1],database=sys.argv[2],milvus_uri=sys.argv[3],provider=HashEmbeddingProvider(32),allow_degraded=True)
s._stage=lambda stage,job: os._exit(23) if stage==sys.argv[4] else None
s.ingest('新版费用200元。',uri='file:///crash.md',title='制度',document_id='policy')
'''
    child=subprocess.run([sys.executable,'-c',script,os.environ.get('RAG_TEST_MONGO_URI','mongodb://127.0.0.1:27028'),
                           store.db.name,vector_uri,stage],capture_output=True,timeout=30)
    assert child.returncode==23, child.stderr.decode()
    # Simulate lease timeout after actual process death, without sleeping 120 seconds.
    store.db.heads.update_many({}, {'$set':{'lease_until':utcnow()-timedelta(seconds=1)}})
    result=store.reconcile();assert not result['errors']
    job=store.db.jobs.find_one({'uri':'file:///crash.md'})
    if stage=='vector_batch_written':
        assert job['status']=='failed' and store.db.chunks.count_documents({'snapshot':job['snapshot_id']})==0
        assert store.snapshot(old['snapshot_id'],collection='library')['status']=='active'
        assert not store.vectors.get(collection_name=job['vector_collection'],ids=job['milvus_ids'])
    else:
        assert job['status']=='completed' and store.snapshot(job['snapshot_id'],collection='library')['status']=='active'
        assert store.snapshot(old['snapshot_id'],collection='library')['status']=='archived'


def test_collection_boundary_and_immutable_source(store):
    result=seed(store)
    assert not store.search('费用',collection='other',mode='bm25')
    with pytest.raises(ValueError):store.snapshot(result['snapshot_id'],collection='other')
    store.db.versions.update_one({'_id':result['snapshot_id']},{'$set':{'content':'tampered'}})
    with pytest.raises(ValueError,match='hash'):store.search('费用',mode='bm25')


@pytest.mark.parametrize('policy',['quota_equal','quota_dense','quota_lexical'])
def test_candidate_quota_keeps_version_scope_source_and_fallback_gates(store,policy):
    old=seed(store)
    cutoff=store.db.versions.find_one({'_id':old['snapshot_id']})['published_at'].isoformat()+'+00:00'
    new=store.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
    def tampering_reranker(query,hits):
        for hit in hits:hit['content']='invented'
        return hits
    store.reranker=tampering_reranker
    trace={}
    hits=store.search('费用',candidate_policy=policy,trace=trace)
    assert hits and all(h['snapshot_id']==new['snapshot_id'] and '200' in h['content'] for h in hits)
    assert trace['policy']==policy and trace['vector_status']=='fallback_bm25'
    assert trace['validated']==[h['chunk_id'] for h in hits]
    assert store.candidate_policy=='rrf'  # per-call experiments cannot change other Workers
    assert not store.search('费用',candidate_policy=policy,collection='other')
    assert '100' in store.search('费用',candidate_policy=policy,as_of=cutoff)[0]['content']
    store.db.versions.update_one({'_id':new['snapshot_id']},{'$set':{'content':'tampered'}})
    with pytest.raises(ValueError,match='hash'):
        store.search('费用',candidate_policy=policy)


def test_capture_time_history_and_archived_reimport(store):
    old=seed(store);cutoff=store.db.versions.find_one({'_id':old['snapshot_id']})['published_at'].isoformat()+'+00:00'
    store.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
    assert '100' in store.search('费用',mode='bm25',as_of=cutoff)[0]['content']
    assert store.ingest('旧版费用为100元。',uri='file:///old.md',title='制度',document_id='policy')['status']=='archived'


def test_metadata_correction_avoids_encoding(store):
    result=seed(store);calls=store.provider.calls
    store.ingest('旧版费用为100元。',uri='file:///old.md',title='修正标题',document_id='policy',metadata={'published_at':'2025-01-01'})
    row=store.snapshot(result['snapshot_id'],collection='library')
    assert row['title']=='修正标题' and json.loads(row['metadata'])['published_at']=='2025-01-01'
    assert store.provider.calls==calls


@pytest.mark.parametrize('suffix,data',[('.pdf',b'not PDF'),('.txt',b'   '),('.exe',b'bad')])
def test_parse_failure_has_job_and_keeps_old(store,tmp_path,suffix,data):
    old=seed(store);path=tmp_path/('bad'+suffix);path.write_bytes(data)
    with pytest.raises(Exception):store.ingest_file(path,document_id='policy')
    assert store.db.jobs.find_one({'phase':'parse_failed'})['status']=='failed'
    assert store.snapshot(old['snapshot_id'],collection='library')['status']=='active'


def test_pdf_pages_and_source_persistence(store,tmp_path):
    from evaluation.knowledge_rag.audit_ingestion import make_pdf
    path=tmp_path/'pages.pdf';make_pdf(path,['Alpha item 100 dollars.','Beta item 200 dollars.'])
    result=store.ingest_file(path);path.unlink()
    hit=store.search('Beta',mode='bm25')[0]
    assert hit['page']==2
    snapshot=store.snapshot(result['snapshot_id'],collection='library')
    assert snapshot['content'][hit['start']:hit['end']]==hit['content']


def test_docx_headings_and_table(store,tmp_path):
    from docx import Document
    doc=Document();doc.add_heading('经费制度',level=1);doc.add_paragraph('预算200万元。')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='部门';table.cell(0,1).text='金额（万元）'
    table.cell(1,0).text='研发';table.cell(1,1).text='100'
    path=tmp_path/'policy.docx';doc.save(path)
    result=store.ingest_file(path)
    snap=store.snapshot(result['snapshot_id'],collection='library')
    assert '# 经费制度' in snap['content'] and '金额（万元）' in snap['content']
    assert store.search('研发',mode='bm25')


def test_large_heading_and_table_source_integrity():
    from deepresearch.knowledge.dual_store import structured_chunks
    heading='# '+'标题'*800
    assert max(p['end']-p['start'] for p in structured_chunks(heading))<=600
    text='| 项目 | 金额（万元） |\n| --- | --- |\n'+'| 研发支出 | 128 |\n'*150
    parts=structured_chunks(text)
    assert len(parts)>1 and all(p['index_prefix'] for p in parts[1:])
    assert all(text[p['start']:p['end']] for p in parts)


def test_final_authority_gate_blocks_archived_source(store):
    from deepresearch.schemas import FetchedDocument
    from deepresearch.evidence import records_from_document
    from deepresearch.knowledge.verification import invalid_library_sources
    old=seed(store);version=store.snapshot(old['snapshot_id'],collection='library')
    doc=FetchedDocument(requested_url=version['uri'],final_url=version['uri'],title='制度',content=version['content'],
                        content_type='text/plain',http_status=200,
                        knowledge_ref={'snapshot_id':old['snapshot_id'],'collection':'library'})
    source,snapshot,*_=records_from_document(doc,'local_knowledge',[])
    assert not invalid_library_sources(store,[source],[snapshot])
    store.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
    assert invalid_library_sources(store,[source],[snapshot])=={source.source_id}
    assert invalid_library_sources(None,[source],[snapshot])=={source.source_id}


def test_real_e5_milvus_vector_search_and_rebuild(store):
    from deepresearch.memory_system.embeddings import ResilientEmbeddingProvider
    store.provider=ResilientEmbeddingProvider(fallback_enabled=False);store.allow_degraded=False
    item=store.ingest('# 报销制度\n\n差旅报销应在十五个工作日内提交。',uri='file:///travel.md',title='差旅报销制度')
    version=store.snapshot(item['snapshot_id'],collection='library')
    assert not item['degraded'] and store.provider.dimensions==384
    hit=store.search('出差结束以后多长时间需要申请费用报销',mode='vector')[0]
    assert hit['snapshot_id']==item['snapshot_id'] and hit['vector_status']=='available'
    name=version['vector_collection'];store.vectors.delete(collection_name=name,ids=[hit['chunk_id']])
    assert not store.search('出差报销',mode='vector')
    assert store.reindex()['reindexed_snapshots']==1
    assert store.search('出差报销',mode='vector')
    after=store.snapshot(item['snapshot_id'],collection='library')
    assert after['content']==version['content'] and after['hash']==version['hash'] and after['status']=='active'


def test_failed_rebuild_keeps_previous_projection(store):
    item=seed(store);before=store.snapshot(item['snapshot_id'],collection='library')
    store._stage=lambda stage,job: (_ for _ in ()).throw(OSError('rebuild interrupted')) if stage=='reindex_batch_written' else None
    with pytest.raises(OSError):store.reindex()
    after=store.snapshot(item['snapshot_id'],collection='library')
    assert after==before and store.search('费用',mode='bm25')
    job=store.db.jobs.find_one({'kind':'index_rebuild'})
    assert job['status']=='failed' and not store.vectors.has_collection(collection_name=job['vector_collection'])


def test_description_index_returns_only_canonical_quotes_and_survives_reader_switch(store):
    from deepresearch.knowledge.document_context import DocumentContextIndex
    from tests.test_knowledge_document_context import context_cache
    content = 'python -m example resume run42'
    store.document_context = DocumentContextIndex(context_cache(content, '断点恢复研究任务的命令。'))
    item = store.ingest(content, uri='file:///resume.md', title='命令')
    version = store.snapshot(item['snapshot_id'], collection='library')
    assert len(version['retrieval_contexts']) == 1 and version['context_manifest_sha256']
    store.document_context = None
    hit = store.search('断点恢复', mode='bm25')[0]
    assert hit['content'] == content and hit['start'] == 0 and hit['end'] == len(content)
    assert '断点恢复' not in json.dumps(hit, ensure_ascii=False)
    assert store.reindex()['reindexed_snapshots'] == 1
    assert not store.search('断点恢复', mode='bm25')
    after = store.snapshot(item['snapshot_id'], collection='library')
    assert after['content'] == version['content'] and after['hash'] == version['hash']
    assert after['retrieval_contexts'] == {} and after['context_manifest_sha256'] is None


def test_failed_description_rebuild_does_not_publish_bm25_notes(store):
    from deepresearch.knowledge.document_context import DocumentContextIndex
    from tests.test_knowledge_document_context import context_cache
    content = 'python -m example resume run42'
    item = store.ingest(content, uri='file:///resume.md', title='命令')
    before = store.snapshot(item['snapshot_id'], collection='library')
    store.document_context = DocumentContextIndex(context_cache(content, '断点恢复研究任务的命令。'))
    def fail(stage, job):
        if stage == 'reindex_batch_written':
            raise OSError('injected rebuild failure')
    store._stage = fail
    with pytest.raises(OSError):
        store.reindex()
    assert store.snapshot(item['snapshot_id'], collection='library') == before
    assert not store.search('断点恢复', mode='bm25')
    assert store.search('resume', mode='bm25')[0]['content'] == content
    vectors = store.vectors.get(collection_name=before['vector_collection'],
        ids=[item['snapshot_id']+':0'], output_fields=['id', 'vector'], consistency_level='Strong')
    assert len(vectors) == 1 and len(vectors[0]['vector']) == store.provider.dimensions


def test_invalid_context_cache_does_not_archive_previous_version(store):
    from deepresearch.knowledge.document_context import DocumentContextIndex
    from tests.test_knowledge_document_context import context_cache
    old = seed(store)
    content = '新版费用200元。'
    cache = context_cache(content)
    cache['chunks']['0']['chunk_sha256'] = '0'*64
    store.document_context = DocumentContextIndex(cache)
    with pytest.raises(ValueError, match='hash mismatch'):
        store.ingest(content, uri='file:///new.md', title='制度', document_id='policy')
    assert store.snapshot(old['snapshot_id'], collection='library')['status'] == 'active'
    assert store.db.jobs.count_documents({'status': 'failed'}) == 1


@pytest.mark.parametrize('rewrite', [False, True])
def test_neighbors_score_only_canonical_source_and_never_expand_returned_evidence(store, rewrite):
    from deepresearch.knowledge.query_rewrite import RewrittenQuerySearch
    content = '恢复上下文。\n\npython -m tool resume run42\n\n输出JSON结果。'
    store.ingest(content, uri='file:///resume.md', title='命令', collection='scoped')
    store.reranker_neighbor_chars = 12
    calls = []
    def score_and_tamper(query, hits):
        calls.append(hits)
        for hit in hits:
            a, b = hit['start'], hit['end']
            assert hit['content'] == content[a:b]
            assert hit['_rerank_neighbors'] == {'before': content[max(0,a-12):a], 'after': content[b:b+12]}
        # Ranking is trusted only as a permutation, never as a new source body.
        return [{**h, 'content': '伪造原文', '_rerank_neighbors': {'before': 'bad', 'after': 'bad'}} for h in hits]
    store.reranker = score_and_tamper
    search = RewrittenQuerySearch(store, {'resume': 'tool resume'}) if rewrite else store
    hits = search.search('resume', mode='bm25', collection='scoped')
    assert calls and hits and all(h['rerank_status'] == 'applied' for h in hits)
    for hit in hits:
        assert hit['content'] == content[hit['start']:hit['end']]
        assert '_rerank_neighbors' not in hit


def test_neighbor_rerank_preserves_source_on_timeout(store):
    content = '前文。\n\nresume command\n\n后文。'
    store.ingest(content, uri='file:///resume.md', title='命令')
    store.reranker_neighbor_chars = 120
    def timeout(query, hits):
        assert any(h['_rerank_neighbors']['before'] for h in hits)
        raise TimeoutError('injected reranker timeout')
    store.reranker = timeout
    hits = store.search('resume', mode='bm25')
    assert hits and all(h['rerank_status'] == 'fallback' for h in hits)
    assert all('_rerank_neighbors' not in h and h['content'] == content[h['start']:h['end']] for h in hits)


def test_stale_vector_replica_does_not_return_archived_facts(store):
    from deepresearch.memory_system.embeddings import ResilientEmbeddingProvider
    store.provider=ResilientEmbeddingProvider(fallback_enabled=False)
    store.allow_degraded=False
    old=seed(store)
    store.ingest('新版费用200元。',uri='file:///new.md',title='制度',document_id='policy')
    # Force the vector replica back to stale active status. Mongo must reject it.
    old_version=store.snapshot(old['snapshot_id'],collection='library')
    rows=store.vectors.get(collection_name=old_version['vector_collection'],ids=[old['snapshot_id']+':0'],
                           output_fields=['id','snapshot','scope','vector','status'])
    store.vectors.upsert(collection_name=old_version['vector_collection'],data=[{**row,'status':'active'} for row in rows])
    store.vectors.flush(collection_name=old_version['vector_collection'])
    hits=store.search('费用',mode='vector')
    assert hits and all(h['vector_status']=='available' for h in hits)
    assert all(h['snapshot_id']!=old['snapshot_id'] and '200' in h['content'] for h in hits)


def test_vector_and_reranker_failure_keep_lexical_results(store,monkeypatch):
    seed(store)
    store.reranker=lambda *args: (_ for _ in ()).throw(TimeoutError('reranker offline'))
    hit=store.search('费用',mode='hybrid')[0]
    assert hit['rerank_status']=='fallback' and '100' in hit['content']
    assert hit['vector_status']=='fallback_bm25'  # explicit hash fixture cannot supply semantics


def test_unavailable_embedding_never_publishes_hash_fallback(store):
    from deepresearch.memory_system.embeddings import HashEmbeddingProvider
    old=seed(store);store.provider=HashEmbeddingProvider(32);store.allow_degraded=False
    with pytest.raises(ValueError,match='semantic'):
        store.ingest('新版200元。',uri='file:///new.md',title='制度',document_id='policy')
    assert store.snapshot(old['snapshot_id'],collection='library')['status']=='active'


def test_precontext_uses_mongo_original_and_preserves_table_header(store):
    from deepresearch.knowledge.context import build_context
    content='| 项目 | 金额（万元） |\n| --- | --- |\n'+'| 研发支出 | 128 |\n'*120
    result=store.ingest(content,uri='file:///table.md',title='经费表')
    pack,_=build_context(store,['研发支出'],limit=3)
    assert pack['status']=='retrieved' and pack['passages']
    assert all(content[p['start']:p['end']]==p['content'] for p in pack['passages'])
    assert all(p['snapshot_id']==result['snapshot_id'] for p in pack['passages'])


def test_compensation_refuses_published_history(store):
    first=seed(store)
    store.ingest('新版200元。',uri='file:///new.md',title='制度',document_id='policy')
    job=store.db.jobs.find_one({'snapshot_id':first['snapshot_id']})
    with pytest.raises(ValueError,match='published'):store._compensate(job)
    assert store.db.chunks.count_documents({'snapshot':first['snapshot_id']})>0


def test_reconcile_respects_live_writer_lease(store):
    seed(store)
    with store._lease('library','policy'):
        store.db.jobs.insert_one({'_id':uuid.uuid4().hex,'status':'pending','snapshot_id':'none',
                                  'collection':'library','document_id':'policy','head_id':store._head_id('library','policy')})
        assert store.reconcile()['busy']==1


def test_retry_rejects_changed_original_upload(store,tmp_path):
    path=tmp_path/'parse.txt';path.write_bytes(b'\xff')
    with pytest.raises(UnicodeDecodeError):store.ingest_file(path)
    job=store.db.jobs.find_one({'status':'failed'});path.write_text('changed source')
    with pytest.raises(ValueError,match='changed'):store.retry(job['_id'])
